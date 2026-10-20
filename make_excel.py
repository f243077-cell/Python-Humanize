"""Builds Test_Cases.xlsx. Every row's input/expected value is read from data/*.csv (external data)."""
import csv, datetime as dt
from pathlib import Path
import humanize
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

D = Path(__file__).parent / "data"
rows = []
for r in csv.DictReader(open(D / "number_cases.csv", encoding="utf-8")):
    out = getattr(humanize, r["function"])(int(r["value"]))
    rows.append([r["test_id"], "humanize.number", r["function"], r["technique"], f"value={r['value']}", "number_cases.csv", r["expected"], out])
for r in csv.DictReader(open(D / "filesize_cases.csv", encoding="utf-8")):
    out = humanize.naturalsize(int(r["value"]), binary=r["binary"] == "true", gnu=r["gnu"] == "true")
    rows.append([r["test_id"], "humanize.filesize", "naturalsize", r["technique"],
                 f"value={r['value']}, binary={r['binary']}, gnu={r['gnu']}", "filesize_cases.csv", r["expected"], out])
for r in csv.DictReader(open(D / "time_cases.csv", encoding="utf-8")):
    out = getattr(humanize, r["function"])(dt.timedelta(seconds=int(r["seconds"])))
    rows.append([r["test_id"], "humanize.time", r["function"], r["technique"], f"seconds={r['seconds']}", "time_cases.csv", r["expected"], out])

wb = Workbook(); ws = wb.active; ws.title = "Test Cases"
F = lambda **k: Font(**{"name": "Arial", "size": 10, **k})
hdr = ["Test ID", "Module", "Function", "Technique", "Test Input", "Data Source (external CSV)",
       "Expected Result", "Actual Result", "Status", "Automated in"]
ws.append(["Assignment 3 - Test Cases (data consumed from CSV files in data/)"])
ws["A1"].font = F(bold=True, size=12)
ws.append(["Legend: Actual Result = output captured from running humanize with the CSV input. Status is a formula (Expected = Actual)."])
ws["A2"].font = F(italic=True)
ws.append([])
ws.append(hdr)
thin = Side(style="thin", color="999999")
for c in ws[4]:
    c.font = F(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, r in enumerate(rows, start=5):
    ws.append(r[:8] + [f'=IF(G{i}=H{i},"Pass","Fail")', "tests/test_csv_driven.py"])
    for c in ws[i]:
        c.font = F(); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in "GH":  # keep values as text
            ws[f"{col}{i}"].number_format = "@"
last = 4 + len(rows)
for col, w in zip("ABCDEFGHIJ", [9, 18, 14, 11, 34, 24, 22, 22, 9, 26]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

s = wb.create_sheet("Summary")
s.append(["Metric", "Value"])
s.append(["Total test cases", f"=COUNTA('Test Cases'!A5:A{last})"])
s.append(["Passed", f"=COUNTIF('Test Cases'!I5:I{last},\"Pass\")"])
s.append(["Failed", f"=COUNTIF('Test Cases'!I5:I{last},\"Fail\")"])
s.append(["Pass rate", "=IF(B2=0,0,B3/B2)"]); s["B5"].number_format = "0.0%"
s.append(["ECP cases", f"=COUNTIF('Test Cases'!D5:D{last},\"ECP\")"])
s.append(["BVA cases", f"=COUNTIF('Test Cases'!D5:D{last},\"BVA\")"])
for row in s.iter_rows():
    for c in row: c.font = F(bold=(c.row == 1))
s.column_dimensions["A"].width = 22; s.column_dimensions["B"].width = 12
wb.save(Path(__file__).parent / "Test_Cases.xlsx")
print(len(rows), "test cases written")
